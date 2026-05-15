from __future__ import annotations

import mailbox
import shutil
from email.message import EmailMessage
from email.utils import format_datetime
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document
from PIL import Image, ImageDraw

from pcad.ingestion.parsers.docx import parse_docx
from pcad.ingestion.parsers.mbox import parse_mbox
from pcad.ingestion.parsers.meeting_md import parse_meeting_md
from pcad.ingestion.parsers.pdf import parse_pdf


def test_mbox_parser_produces_messages_with_thread_metadata(tmp_path: Path) -> None:
    path = tmp_path / "emails.mbox"
    _write_mbox(path)

    emails = parse_mbox(path)

    assert len(emails) == 3
    references = {email.in_reply_to for email in emails if email.in_reply_to}
    assert references == {"root@example.test"}


def test_pdf_parser_extracts_text_layer(tmp_path: Path) -> None:
    path = tmp_path / "proposal.pdf"
    _write_minimal_text_pdf(path, "Security review approved for the pilot proposal.")

    parsed = parse_pdf(path, ocr_fallback=False)

    assert parsed.page_count == 1
    assert "Security review approved" in parsed.text_per_page[0]
    assert parsed.extraction_method == "plain_text"


@pytest.mark.requires_tesseract
def test_pdf_parser_ocr_fallback_for_image_pdf(tmp_path: Path) -> None:
    if not shutil.which("tesseract") or not shutil.which("pdftoppm"):
        pytest.skip("Tesseract or Poppler is not installed")
    path = tmp_path / "scanned.pdf"
    image = Image.new("RGB", (900, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 120), "Scanned NDA approval text", fill="black")
    image.save(path, "PDF")

    parsed = parse_pdf(path, ocr_fallback=True)

    assert parsed.extraction_method == "ocr"
    assert "NDA" in " ".join(parsed.text_per_page)


def test_docx_parser_preserves_headings(tmp_path: Path) -> None:
    path = tmp_path / "account_plan.docx"
    document = Document()
    document.add_heading("Account Plan", level=1)
    document.add_heading("Risks", level=2)
    document.add_paragraph("Procurement needs a follow-up.")
    document.save(path)

    parsed = parse_docx(path)

    assert "# Account Plan" in parsed.text
    assert "## Risks" in parsed.text


def test_docx_parser_promotes_title_and_subtitle(tmp_path: Path) -> None:
    path = tmp_path / "account_plan.docx"
    document = Document()
    document.add_heading("Account Plan", level=0)
    document.add_paragraph("Private memory pilot", style="Subtitle")
    document.add_paragraph("Procurement needs a follow-up.")
    document.save(path)

    parsed = parse_docx(path)

    assert "# Account Plan" in parsed.text
    assert "## Private memory pilot" in parsed.text


def test_meeting_markdown_parser_extracts_header_and_turns(tmp_path: Path) -> None:
    path = tmp_path / "meeting.md"
    path.write_text(
        """# Meeting: Pilot review

**Date:** 2026-05-01
**Attendees:** Daniel Panea, Jordan Lee

---

**Daniel:** Thanks for joining.
**Jordan:** The main concern is security approval.
""",
        encoding="utf-8",
    )

    parsed = parse_meeting_md(path)

    assert parsed.meeting_title == "Meeting: Pilot review"
    assert parsed.meeting_date and parsed.meeting_date.isoformat() == "2026-05-01"
    assert parsed.attendees == ["Daniel Panea", "Jordan Lee"]
    assert len(parsed.turns) == 2


def _write_mbox(path: Path) -> None:
    box = mailbox.mbox(path)
    box.lock()
    try:
        root_id = "root@example.test"
        box.add(_message(root_id, "Pilot plan", "daniel@example.test", ["jordan@example.test"], "Can we approve the next step?"))
        box.add(
            _message(
                "reply@example.test",
                "Re: Pilot plan",
                "jordan@example.test",
                ["daniel@example.test"],
                "Approval depends on procurement.",
                in_reply_to=root_id,
                references=[root_id],
            )
        )
        box.add(_message("other@example.test", "Separate topic", "avery@example.test", ["daniel@example.test"], "This is a new thread."))
        box.flush()
    finally:
        box.unlock()
        box.close()


def _message(
    message_id: str,
    subject: str,
    from_addr: str,
    to_addrs: list[str],
    body: str,
    *,
    in_reply_to: str | None = None,
    references: list[str] | None = None,
) -> EmailMessage:
    message = EmailMessage()
    message["Message-ID"] = f"<{message_id}>"
    message["Subject"] = subject
    message["From"] = from_addr
    message["To"] = ", ".join(to_addrs)
    message["Date"] = format_datetime(datetime(2026, 5, 1, 10, 0, tzinfo=timezone.utc))
    if in_reply_to:
        message["In-Reply-To"] = f"<{in_reply_to}>"
    if references:
        message["References"] = " ".join(f"<{item}>" for item in references)
    message.set_content(body)
    return message


def _write_minimal_text_pdf(path: Path, text: str) -> None:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode("ascii"))
        content.extend(obj)
        content.extend(b"\nendobj\n")
    xref_at = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(bytes(content))
